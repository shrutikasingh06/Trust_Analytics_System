"""Rebuild synopsis front matter: LOF, LOT, TOC, roman/arabic paging, 10 refs."""

from __future__ import annotations

import copy
import re
import shutil
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips

SRC = Path(r"D:\Trust_Analytics_System\Shrutika_Project2_Synopsis.docx")
WORK = Path(r"D:\Trust_Analytics_System\Shrutika_Project2_Synopsis_work.docx")
FINAL_DOWNLOADS = Path(r"C:\Users\HP\Downloads\Shrutika_Project2_Synopsis_Final.docx")
FINAL_WS = Path(r"D:\Trust_Analytics_System\Shrutika_Project2_Synopsis_Final.docx")
BACKUP = Path(r"C:\Users\HP\Downloads\Shrutika_Project2_Synopsis_backup.docx")

FIGURES = [
    ("Fig. 1.1", "Categories of Suspicious Review Indicators", "fig_1_1"),
    ("Fig. 1.2", "The Trust Gap in Online Reviews", "fig_1_2"),
    ("Fig. 3.1", "Proposed System Architecture", "fig_3_1"),
    ("Fig. 3.2", "Module Interaction and Data Flow", "fig_3_2"),
    ("Fig. 4.1", "Overall Methodology Flowchart", "fig_4_1"),
    ("Fig. 4.2", "Trust Score Generation Flowchart", "fig_4_2"),
]

TABLES = [
    ("Table 2.1", "Comparison of Existing Approaches and Proposed System", "tbl_2_1"),
    ("Table 2.2", "Technique Comparison by Practical Criteria", "tbl_2_2"),
    ("Table 2.3", "Summary of Key Studies Reviewed", "tbl_2_3"),
    ("Table 3.1", "Hardware Requirements", "tbl_3_1"),
    ("Table 3.2", "Software Requirements", "tbl_3_2"),
    ("Table 3.3", "Expected Dataset Field Categories", "tbl_3_3"),
    ("Table 4.1", "Proposed Data-Cleaning Actions by Issue Type", "tbl_4_1"),
    ("Table 4.2", "Major Analytical Features", "tbl_4_2"),
    ("Table 4.3", "Illustrative (Hypothetical) Trust Score Composition", "tbl_4_3"),
    ("Table 4.4", "Approximate Planned Duration per Implementation Stage", "tbl_4_4"),
    ("Table 4.5", "Project Implementation Schedule", "tbl_4_5"),
    ("Table 4.6", "Implementation Risk Analysis", "tbl_4_6"),
    ("Table 5.1", "Objectives-to-Coverage Mapping", "tbl_5_1"),
]

TOC_ITEMS = [
    ("Approval Sheet", "h_approval", 0),
    ("Abstract", "h_abstract", 0),
    ("List of Figures", "h_lof", 0),
    ("List of Tables", "h_lot", 0),
    ("Table of Contents", "h_toc", 0),
    ("Chapter 1: Introduction", "h_ch1", 0),
    ("1.1  Overview of E-Commerce", "h_1_1", 1),
    ("1.2  Importance of Online Reviews", "h_1_2", 1),
    ("1.3  Problem of Fake and Suspicious Reviews", "h_1_3", 1),
    ("1.4  Need for E-Commerce Trust Analysis", "h_1_4", 1),
    ("1.5  Problem Statement", "h_1_5", 1),
    ("1.6  Objectives", "h_1_6", 1),
    ("1.7  Scope of the Proposed System", "h_1_7", 1),
    ("1.8  Expected Outcome", "h_1_8", 1),
    ("1.9  Motivation for the Project", "h_1_9", 1),
    ("1.10  Organisation of the Synopsis", "h_1_10", 1),
    ("1.11  Illustrative Scenarios", "h_1_11", 1),
    ("1.12  Key Terms and Definitions", "h_1_12", 1),
    ("1.13  Comparison with Traditional (Offline) Trust Signals", "h_1_13", 1),
    ("Chapter 2: Literature Review", "h_ch2", 0),
    ("2.1  E-Commerce Review Systems", "h_2_1", 1),
    ("2.2  Fake Review Detection", "h_2_2", 1),
    ("2.3  Review Text Analysis", "h_2_3", 1),
    ("2.4  Reviewer Behaviour Analysis", "h_2_4", 1),
    ("2.5  Rating-Based Analysis", "h_2_5", 1),
    ("2.6  Behavioural and Temporal Analysis", "h_2_6", 1),
    ("2.7  Existing Approaches", "h_2_7", 1),
    ("2.8  Limitations of Existing Approaches", "h_2_8", 1),
    ("2.9  Research Gap", "h_2_9", 1),
    ("2.10  Proposed Improvement", "h_2_10", 1),
    ("2.11  Summary of Literature Review", "h_2_11", 1),
    ("2.12  Comparative Analysis of Detection Techniques", "h_2_12", 1),
    ("2.13  Ethical and Privacy Considerations", "h_2_13", 1),
    ("2.14  Summary Table of Key Studies", "h_2_14", 1),
    ("Chapter 3: Proposed System", "h_ch3", 0),
    ("3.1  System Architecture / Block Diagram", "h_3_1", 1),
    ("3.2  Module-wise Description", "h_3_2", 1),
    ("3.3  Hardware and Software Requirements", "h_3_3", 1),
    ("3.4  Data Dictionary", "h_3_4", 1),
    ("3.5  Security Considerations", "h_3_5", 1),
    ("3.6  Deployment Considerations (Proposed)", "h_3_6", 1),
    ("3.7  Alternative Design Considered", "h_3_7", 1),
    ("3.8  Assumptions and Constraints", "h_3_8", 1),
    ("3.9  Advantages and Disadvantages", "h_3_9", 1),
    ("Chapter 4: Methodology", "h_ch4", 0),
    ("4.1  Techniques / Methods", "h_4_1", 1),
    ("4.2  Algorithms", "h_4_2", 1),
    ("4.3  Flowcharts", "h_4_3", 1),
    ("4.4  Plan of Implementation", "h_4_4", 1),
    ("4.5  Schedule of Work", "h_4_5", 1),
    ("4.6  Proposed Evaluation Strategy", "h_4_6", 1),
    ("4.7  Risk Analysis and Mitigation", "h_4_7", 1),
    ("4.8  Proposed Testing Strategy", "h_4_8", 1),
    ("4.9  Expected Deliverables", "h_4_9", 1),
    ("4.10  Tools Justification", "h_4_10", 1),
    ("Chapter 5: Conclusion", "h_ch5", 0),
    ("5.1  Conclusion", "h_5_1", 1),
    ("5.2  Expected Benefits", "h_5_2", 1),
    ("5.3  Limitations of the Proposed Study", "h_5_3", 1),
    ("5.4  Future Scope", "h_5_4", 1),
    ("5.5  Summary", "h_5_5", 1),
    ("5.6  Personal Learning Outcomes", "h_5_6", 1),
    ("5.7  Mapping of Objectives to Proposed Coverage", "h_5_7", 1),
    ("Chapter 6 -- References", "h_ch6", 0),
]

HEADING_BOOKMARKS = [
    ("APPROVAL SHEET", "h_approval"),
    ("ABSTRACT", "h_abstract"),
    ("LIST OF FIGURES", "h_lof"),
    ("LIST OF TABLES", "h_lot"),
    ("TABLE OF CONTENTS", "h_toc"),
    ("Chapter 1: Introduction", "h_ch1"),
    ("1.1 Overview of E-Commerce", "h_1_1"),
    ("1.2 Importance of Online Reviews", "h_1_2"),
    ("1.3 Problem of Fake and Suspicious Reviews", "h_1_3"),
    ("1.4 Need for E-Commerce Trust Analysis", "h_1_4"),
    ("1.5 Problem Statement", "h_1_5"),
    ("1.6 Objectives", "h_1_6"),
    ("1.7 Scope of the Proposed System", "h_1_7"),
    ("1.8 Expected Outcome", "h_1_8"),
    ("1.9 Motivation for the Project", "h_1_9"),
    ("1.10 Organisation of the Synopsis", "h_1_10"),
    ("1.11 Illustrative Scenarios", "h_1_11"),
    ("1.12 Key Terms and Definitions", "h_1_12"),
    ("1.13 Comparison with Traditional (Offline) Trust Signals", "h_1_13"),
    ("Chapter 2: Literature Review", "h_ch2"),
    ("2.1 E-Commerce Review Systems", "h_2_1"),
    ("2.2 Fake Review Detection", "h_2_2"),
    ("2.3 Review Text Analysis", "h_2_3"),
    ("2.4 Reviewer Behaviour Analysis", "h_2_4"),
    ("2.5 Rating-Based Analysis", "h_2_5"),
    ("2.6 Behavioural and Temporal Analysis", "h_2_6"),
    ("2.7 Existing Approaches", "h_2_7"),
    ("2.8 Limitations of Existing Approaches", "h_2_8"),
    ("2.9 Research Gap", "h_2_9"),
    ("2.10 Proposed Improvement", "h_2_10"),
    ("2.11 Summary of Literature Review", "h_2_11"),
    ("2.12 Comparative Analysis of Detection Techniques", "h_2_12"),
    ("2.13 Ethical and Privacy Considerations", "h_2_13"),
    ("2.14 Summary Table of Key Studies", "h_2_14"),
    ("Chapter 3: Proposed System", "h_ch3"),
    ("3.1 System Architecture / Block Diagram", "h_3_1"),
    ("3.2 Module-wise Description", "h_3_2"),
    ("3.3 Hardware and Software Requirements", "h_3_3"),
    ("3.4 Data Dictionary", "h_3_4"),
    ("3.5 Security Considerations", "h_3_5"),
    ("3.6 Deployment Considerations (Proposed)", "h_3_6"),
    ("3.7 Alternative Design Considered", "h_3_7"),
    ("3.8 Assumptions and Constraints", "h_3_8"),
    ("3.9 Advantages and Disadvantages", "h_3_9"),
    ("Chapter 4: Methodology", "h_ch4"),
    ("4.1 Techniques / Methods", "h_4_1"),
    ("4.2 Algorithms", "h_4_2"),
    ("4.3 Flowcharts", "h_4_3"),
    ("4.4 Plan of Implementation", "h_4_4"),
    ("4.5 Schedule of Work", "h_4_5"),
    ("4.6 Proposed Evaluation Strategy", "h_4_6"),
    ("4.7 Risk Analysis and Mitigation", "h_4_7"),
    ("4.8 Proposed Testing Strategy", "h_4_8"),
    ("4.9 Expected Deliverables", "h_4_9"),
    ("4.10 Tools Justification", "h_4_10"),
    ("Chapter 5: Conclusion", "h_ch5"),
    ("5.1 Conclusion", "h_5_1"),
    ("5.2 Expected Benefits", "h_5_2"),
    ("5.3 Limitations of the Proposed Study", "h_5_3"),
    ("5.4 Future Scope", "h_5_4"),
    ("5.5 Summary", "h_5_5"),
    ("5.6 Personal Learning Outcomes", "h_5_6"),
    ("5.7 Mapping of Objectives to Proposed Coverage", "h_5_7"),
    ("Chapter 6 -- References", "h_ch6"),
]

REFS = [
    [
        "1. Jindal, N., & Liu, B. (2008). Opinion Spam and Analysis.",
        "Proceedings of the 2008 International Conference on Web Search and Data Mining (WSDM).",
        "Foundational study establishing patterns of duplicate and deceptive review content in e-commerce platforms.",
    ],
    [
        "2. Ott, M., Choi, Y., Cardie, C., & Hancock, J. T. (2011). Finding Deceptive Opinion Spam by Any Stretch of the Imagination.",
        "Proceedings of the 49th Annual Meeting of the Association for Computational Linguistics (ACL).",
        "Demonstrates the use of linguistic and text-based features to distinguish deceptive reviews from genuine ones.",
    ],
    [
        "3. Mukherjee, A., Venkataraman, V., Liu, B., & Glance, N. (2013). What Yelp Fake Review Filter Might Be Doing?",
        "Proceedings of the International AAAI Conference on Web and Social Media (ICWSM).",
        "Analyses reviewer-behaviour signals used by real-world review-filtering systems.",
    ],
    [
        "4. Lim, E.-P., Nguyen, V.-A., Jindal, N., Liu, B., & Lauw, H. W. (2010). Detecting Product Review Spammers Using Rating Behaviors.",
        "Proceedings of the 19th ACM International Conference on Information and Knowledge Management (CIKM).",
        "Proposes reviewer-centred behavioural scoring indicators, including rating deviation and early-review targeting.",
    ],
    [
        "5. Feng, S., Xing, L., Gogar, A., & Choi, Y. (2012). Distributional Footprints of Deceptive Product Reviews.",
        "Proceedings of the International AAAI Conference on Web and Social Media (ICWSM).",
        "Studies corpus-level statistical differences between genuine and deceptive review populations.",
    ],
    [
        "6. Wang, G., Xie, S., Liu, B., & Yu, P. S. (2011). Review Graph Based Online Store Review Spammer Detection.",
        "Proceedings of the 2011 IEEE International Conference on Data Mining (ICDM).",
        "Models reviewers, reviews, and stores as a graph to identify spammer accounts by structural position.",
    ],
    [
        "7. Rayana, S., & Akoglu, L. (2015). Collective Opinion Spam Detection: Bridging Review Networks and Metadata.",
        "Proceedings of the 21st ACM SIGKDD International Conference on Knowledge Discovery and Data Mining (KDD).",
        "Combines reviewer-product network structure with rating and temporal metadata to detect collectively acting spam groups.",
    ],
    [
        "8. Liu, F. T., Ting, K. M., & Zhou, Z.-H. (2008). Isolation Forest.",
        "Proceedings of the 2008 IEEE International Conference on Data Mining (ICDM).",
        "Original paper describing the Isolation Forest anomaly-detection algorithm proposed for behavioural outlier analysis.",
    ],
    [
        "9. Manning, C. D., Raghavan, P., & Schutze, H. (2008). Introduction to Information Retrieval.",
        "Cambridge University Press.",
        "Standard reference textbook covering TF-IDF weighting and vector-space text-similarity models.",
    ],
    [
        "10. McAuley Lab, University of California San Diego — Amazon Review Data.",
        "https://cseweb.ucsd.edu/~jmcauley/datasets/amazon_v2/",
        "Large-scale public e-commerce review dataset resource referenced for planning the dataset characteristics used in this proposal.",
    ],
]


def set_run_font(run, name="Times New Roman", size=12, bold=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)
    rFonts.set(qn("w:cs"), name)


def page_break_before(paragraph, enabled=True):
    pPr = paragraph._p.get_or_add_pPr()
    for el in pPr.findall(qn("w:pageBreakBefore")):
        pPr.remove(el)
    if enabled:
        pb = OxmlElement("w:pageBreakBefore")
        pPr.append(pb)


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def _run():
        r = paragraph.add_run()
        r.font.size = Pt(11)
        return r._r

    r1 = _run()
    b = OxmlElement("w:fldChar")
    b.set(qn("w:fldCharType"), "begin")
    r1.append(b)

    r2 = _run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    r2.append(instr)

    r3 = _run()
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    r3.append(sep)

    r4 = paragraph.add_run("1")
    r4.font.size = Pt(11)

    r5 = _run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    r5.append(end)


def set_pg_num(section, fmt: str | None, start: int | None):
    sectPr = section._sectPr
    for el in list(sectPr):
        if el.tag == qn("w:pgNumType"):
            sectPr.remove(el)
    pg = OxmlElement("w:pgNumType")
    if fmt:
        pg.set(qn("w:fmt"), fmt)
    if start is not None:
        pg.set(qn("w:start"), str(start))
    # Place before pgSz/pgMar if possible; appending is accepted by Word.
    sectPr.append(pg)


def bookmark_paragraph(paragraph, name: str, bid: int):
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bid))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bid))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def find_para(doc, text: str):
    target = re.sub(r"\s+", " ", text).strip()
    for p in doc.paragraphs:
        if re.sub(r"\s+", " ", p.text).strip() == target:
            return p
    return None


def delete_element(el):
    parent = el.getparent()
    if parent is not None:
        parent.remove(el)


def set_cell_text(cell, text, bold=False, size=11, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(2)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)


def shade_cell(cell, fill="1F4E79"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_color(cell, rgb="FFFFFF"):
    color = RGBColor(int(rgb[0:2], 16), int(rgb[2:4], 16), int(rgb[4:6], 16))
    for p in cell.paragraphs:
        for r in p.runs:
            r.font.color.rgb = color


def make_three_col_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1 + len(rows), cols=3)
    try:
        table.style = "Table Grid"
    except Exception:
        if doc.tables:
            try:
                table.style = doc.tables[0].style
            except Exception:
                pass
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for i, w in enumerate(widths):
        for cell in table.columns[i].cells:
            cell.width = w
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, bold=True, size=11, align=WD_ALIGN_PARAGRAPH.CENTER)
        shade_cell(table.rows[0].cells[i], "1F4E79")
        set_cell_color(table.rows[0].cells[i], "FFFFFF")
    for r_i, row in enumerate(rows, start=1):
        align_last = WD_ALIGN_PARAGRAPH.CENTER
        set_cell_text(table.rows[r_i].cells[0], row[0], bold=True, size=11, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(table.rows[r_i].cells[1], row[1], size=11)
        set_cell_text(table.rows[r_i].cells[2], row[2], size=11, align=align_last)
    return table


def insert_before(anchor_elm, new_elm):
    anchor_elm.addprevious(new_elm)


def add_heading_para(doc, text, bookmark, bid, page_break=False):
    p = doc.add_paragraph()
    p.style = doc.styles["Heading 1"]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_run_font(run, size=16, bold=True)
    if page_break:
        page_break_before(p, True)
    bookmark_paragraph(p, bookmark, bid)
    return p


def replace_refs(doc):
    ch6 = find_para(doc, "Chapter 6 -- References")
    if ch6 is None:
        raise RuntimeError("Chapter 6 heading not found")
    # Delete everything after the heading.
    nxt = ch6._p.getnext()
    while nxt is not None:
        following = nxt.getnext()
        tag = nxt.tag
        if tag == qn("w:sectPr"):
            break
        delete_element(nxt)
        nxt = following
    # Append refs just after heading (before sectPr).
    anchor = ch6._p
    for block in REFS:
        for i, line in enumerate(block):
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(0 if i < 2 else 8)
            p.paragraph_format.space_before = Pt(6 if i == 0 else 0)
            run = p.add_run(line)
            set_run_font(run, size=12, bold=(i == 0))
            el = p._p
            el.getparent().remove(el)
            anchor.addnext(el)
            anchor = el
        # blank spacer already via space_after


def insert_section_before_chapter1(doc):
    ch1 = find_para(doc, "Chapter 1: Introduction")
    if ch1 is None:
        raise RuntimeError("Chapter 1 heading not found")
    page_break_before(ch1, False)

    last_sect = doc.sections[-1]._sectPr
    new_p = OxmlElement("w:p")
    pPr = OxmlElement("w:pPr")
    sect = copy.deepcopy(last_sect)
    # Front-matter continuation: lowercase roman, do not restart.
    for el in list(sect):
        if el.tag == qn("w:pgNumType"):
            sect.remove(el)
    pg = OxmlElement("w:pgNumType")
    pg.set(qn("w:fmt"), "lowerRoman")
    sect.append(pg)
    # Ensure next-page section break.
    for el in list(sect):
        if el.tag == qn("w:type"):
            sect.remove(el)
    st = OxmlElement("w:type")
    st.set(qn("w:val"), "nextPage")
    sect.insert(0, st)
    pPr.append(sect)
    new_p.append(pPr)
    ch1._p.addprevious(new_p)

    # Body section: Arabic, restart at 1.
    set_pg_num(doc.sections[-1], "decimal", 1)


def style_list_table_move(doc, table, anchor_elm):
    tbl = table._tbl
    tbl.getparent().remove(tbl)
    insert_before(anchor_elm, tbl)
    return tbl


def rebuild_front_lists_and_toc(doc):
    toc = find_para(doc, "TABLE OF CONTENTS")
    if toc is None:
        raise RuntimeError("TOC heading not found")
    bookmark_paragraph(toc, "h_toc", 200)

    # Remove existing TOC entries (everything between TOC heading and Chapter 1, exclusive).
    ch1 = find_para(doc, "Chapter 1: Introduction")
    el = toc._p.getnext()
    while el is not None and el is not ch1._p:
        nxt = el.getnext()
        # stop if we hit the new section para later; at this stage section not inserted yet
        delete_element(el)
        el = nxt

    # Create LOF/LOT at end then move before TOC.
    lof_p = add_heading_para(doc, "LIST OF FIGURES", "h_lof", 201, page_break=True)
    lof_note = doc.add_paragraph()
    r = lof_note.add_run(
        "The following figures appear in this synopsis. Page numbers follow the printed pagination "
        "(roman numerals for front matter; Arabic numerals from Chapter 1)."
    )
    set_run_font(r, size=11)
    lof_rows = [(n, d, "—") for n, d, _ in FIGURES]
    lof_tbl = make_three_col_table(
        doc,
        ["Fig. No.", "Description", "Page No."],
        lof_rows,
        [Inches(1.2), Inches(4.6), Inches(1.0)],
    )

    lot_p = add_heading_para(doc, "LIST OF TABLES", "h_lot", 202, page_break=True)
    lot_note = doc.add_paragraph()
    r = lot_note.add_run(
        "The following tables appear in this synopsis. Page numbers follow the printed pagination "
        "(roman numerals for front matter; Arabic numerals from Chapter 1)."
    )
    set_run_font(r, size=11)
    lot_rows = [(n, d, "—") for n, d, _ in TABLES]
    lot_tbl = make_three_col_table(
        doc,
        ["Table No.", "Description", "Page No."],
        lot_rows,
        [Inches(1.2), Inches(4.6), Inches(1.0)],
    )

    # Placeholder TOC entries so layout is stable before page lookup.
    toc_paras = []
    for title, _bm, level in TOC_ITEMS:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.space_before = Pt(4 if level == 0 else 0)
        tab_pos = Twips(9360)  # ~6.5"
        p.paragraph_format.tab_stops.add_tab_stop(
            tab_pos, WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS
        )
        p.paragraph_format.left_indent = Inches(0.3 * level)
        run = p.add_run(f"{title}\t—")
        set_run_font(run, size=12, bold=(level == 0))
        toc_paras.append(p)

    # Move newly created blocks to sit before TOC: LOF heading, note, table, LOT..., then TOC entries after TOC heading.
    # Current end-of-body order: lof_p, lof_note, lof_tbl, lot_p, lot_note, lot_tbl, toc_paras...
    chunk_elms = [
        lof_p._p,
        lof_note._p,
        lof_tbl._tbl,
        lot_p._p,
        lot_note._p,
        lot_tbl._tbl,
    ]
    for elm in chunk_elms:
        parent = elm.getparent()
        parent.remove(elm)
        insert_before(toc._p, elm)

    # TOC entries currently at end; move after TOC heading (before Chapter 1).
    ch1 = find_para(doc, "Chapter 1: Introduction")
    insert_point = toc._p
    for p in toc_paras:
        elm = p._p
        elm.getparent().remove(elm)
        insert_point.addnext(elm)
        insert_point = elm

    return lof_tbl, lot_tbl, toc_paras


def bookmark_existing(doc):
    bid = 1
    used = set()
    mapping = {re.sub(r"\s+", " ", t).strip(): bm for t, bm in HEADING_BOOKMARKS}
    for p in doc.paragraphs:
        style = p.style.name if p.style else ""
        if not style.startswith("Heading"):
            continue
        key = re.sub(r"\s+", " ", p.text).strip()
        if key in mapping and mapping[key] not in used:
            if p._p.find(qn("w:bookmarkStart")) is None:
                bookmark_paragraph(p, mapping[key], bid)
                bid += 1
            used.add(mapping[key])
    cap_map = {}
    for num, desc, bm in FIGURES:
        cap_map[f"{num}: {desc}"] = bm
    for num, desc, bm in TABLES:
        cap_map[f"{num}: {desc}"] = bm
    for p in doc.paragraphs:
        key = re.sub(r"\s+", " ", p.text).strip()
        if key in cap_map:
            bookmark_paragraph(p, cap_map[key], bid)
            bid += 1
    return bid


def configure_title_footer(doc):
    # Section 0: title page — roman, start at i.
    set_pg_num(doc.sections[0], "lowerRoman", 1)
    footer = doc.sections[0].footer
    footer.is_linked_to_previous = False
    # Clear existing empty para and add PAGE.
    for p in footer.paragraphs:
        p.clear()
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    add_page_field(p)

    # Section 1 (approval through TOC after we insert break): continue roman.
    if len(doc.sections) >= 2:
        set_pg_num(doc.sections[1], "lowerRoman", None)
        doc.sections[1].footer.is_linked_to_previous = False
        # Keep existing PAGE field.
    if len(doc.sections) >= 3:
        set_pg_num(doc.sections[-1], "decimal", 1)
        doc.sections[-1].footer.is_linked_to_previous = False
        # Keep footer PAGE field; numbering format is section-level.


def to_roman(n: int) -> str:
    vals = [
        (1000, "m"),
        (900, "cm"),
        (500, "d"),
        (400, "cd"),
        (100, "c"),
        (90, "xc"),
        (50, "l"),
        (40, "xl"),
        (10, "x"),
        (9, "ix"),
        (5, "v"),
        (4, "iv"),
        (1, "i"),
    ]
    out = []
    for v, s in vals:
        while n >= v:
            out.append(s)
            n -= v
    return "".join(out)


def fill_pages_with_word(work_path: Path, pages: dict[str, str]):
    doc = Document(str(work_path))

    # Fill LOF
    # Find table whose header is Fig. No.
    for table in doc.tables:
        hdr = table.cell(0, 0).text.strip()
        if hdr in ("Fig. No.", "Fig No.", "Fig. No"):
            for i, (num, desc, bm) in enumerate(FIGURES, start=1):
                if i < len(table.rows):
                    set_cell_text(table.rows[i].cells[2], pages.get(bm, "—"), size=11, align=WD_ALIGN_PARAGRAPH.CENTER)
        elif hdr in ("Table No.", "Table No"):
            for i, (num, desc, bm) in enumerate(TABLES, start=1):
                if i < len(table.rows):
                    set_cell_text(table.rows[i].cells[2], pages.get(bm, "—"), size=11, align=WD_ALIGN_PARAGRAPH.CENTER)

    # Fill TOC paragraphs between TOC heading and Chapter 1 section/heading.
    toc = find_para(doc, "TABLE OF CONTENTS")
    p = toc._p.getnext()
    idx = 0
    while p is not None and idx < len(TOC_ITEMS):
        if p.tag == qn("w:p"):
            texts = [t.text or "" for t in p.iter(qn("w:t"))]
            joined = "".join(texts).strip()
            if joined:
                title, bm, level = TOC_ITEMS[idx]
                page = pages.get(bm, "—")
                # rewrite runs
                for child in list(p):
                    if child.tag != qn("w:pPr"):
                        p.remove(child)
                r = OxmlElement("w:r")
                rPr = OxmlElement("w:rPr")
                sz = OxmlElement("w:sz")
                sz.set(qn("w:val"), "24")
                rPr.append(sz)
                if level == 0:
                    b = OxmlElement("w:b")
                    rPr.append(b)
                rFonts = OxmlElement("w:rFonts")
                rFonts.set(qn("w:ascii"), "Times New Roman")
                rFonts.set(qn("w:hAnsi"), "Times New Roman")
                rPr.append(rFonts)
                r.append(rPr)
                t = OxmlElement("w:t")
                t.set(qn("xml:space"), "preserve")
                t.text = f"{title}\t{page}"
                r.append(t)
                p.append(r)
                idx += 1
        if p.find(qn("w:pPr")) is not None and p.find(qn("w:pPr")).find(qn("w:sectPr")) is not None:
            break
        p = p.getnext()

    doc.save(str(work_path))


def collect_pages_via_word(path: Path) -> dict[str, str]:
    import win32com.client

    # In this Word build, Range.Information(1) is the printed (section) page number.
    wd_printed_page = 1
    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False
    word.DisplayAlerts = 0
    pages: dict[str, str] = {}
    try:
        d = word.Documents.Open(str(path), ReadOnly=False)
        try:
            d.Fields.Update()
        except Exception:
            pass
        d.Repaginate()
        names = [bm for _, bm in HEADING_BOOKMARKS]
        names += [bm for _, _, bm in FIGURES]
        names += [bm for _, _, bm in TABLES]
        front = {"h_approval", "h_abstract", "h_lof", "h_lot", "h_toc"}
        for name in names:
            try:
                rng = d.Bookmarks(name).Range
                n = int(rng.Information(wd_printed_page))
                sec = int(rng.Sections(1).Index)
                if name in front or sec <= 2:
                    pages[name] = to_roman(n)
                else:
                    pages[name] = str(n)
            except Exception:
                pages[name] = "—"
        d.Close(False)
    finally:
        word.Quit()
    return pages


def main():
    orig = Path(r"C:\Users\HP\Downloads\Shrutika_Project2_Synopsis.docx")
    source = BACKUP if BACKUP.exists() else (SRC if SRC.exists() else orig)
    if orig.exists() and not BACKUP.exists():
        shutil.copy2(orig, BACKUP)
        source = BACKUP
    shutil.copy2(source, WORK)

    doc = Document(str(WORK))
    replace_refs(doc)
    rebuild_front_lists_and_toc(doc)
    insert_section_before_chapter1(doc)
    bookmark_existing(doc)
    configure_title_footer(doc)
    doc.save(str(WORK))

    pages = collect_pages_via_word(WORK)
    print("PAGES", pages)
    fill_pages_with_word(WORK, pages)

    # Final Word open to update fields and save copies.
    import win32com.client

    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False
    word.DisplayAlerts = 0
    try:
        d = word.Documents.Open(str(WORK))
        try:
            d.Fields.Update()
        except Exception:
            pass
        d.Repaginate()
        d.SaveAs2(str(FINAL_DOWNLOADS), FileFormat=16)
        d.Close(False)
    finally:
        word.Quit()
    shutil.copy2(FINAL_DOWNLOADS, FINAL_WS)
    shutil.copy2(FINAL_DOWNLOADS, orig)
    print("Wrote", FINAL_DOWNLOADS)
    print("Wrote", FINAL_WS)
    print("Updated", orig)
    print("Backup", BACKUP)


if __name__ == "__main__":
    main()
